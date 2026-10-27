"""Draft Word report builder (python-docx). Every number inserted comes from the computed results."""
import datetime as dt
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .references import REFS


class Report:
    def __init__(self, title, subtitle=""):
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(11)
        self.doc.add_heading(title, 0)
        if subtitle:
            self.doc.add_paragraph(subtitle).runs[0].italic = True
        self.doc.add_paragraph(f"Generated {dt.date.today():%d %B %Y} | AOI: Pakistan")
        self._f = 0
        self._t = 0

    def h(self, text, level=1):
        self.doc.add_heading(text, level)

    def p(self, text):
        para = self.doc.add_paragraph(text)
        para.paragraph_format.space_after = Pt(6)
        return para

    def bullets(self, items):
        for i in items:
            self.doc.add_paragraph(str(i), style="List Bullet")

    def fig(self, path, caption, width=6.0):
        if not Path(path).exists():
            return
        self._f += 1
        self.doc.add_picture(str(path), width=Inches(width))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = self.doc.add_paragraph(f"Figure {self._f}. {caption}")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].italic = True

    def table(self, df: pd.DataFrame, caption, max_rows=25):
        self._t += 1
        self.doc.add_paragraph(f"Table {self._t}. {caption}").runs[0].bold = True
        df = df.head(max_rows)
        t = self.doc.add_table(rows=1, cols=len(df.columns))
        t.style = "Light Grid Accent 1"
        for i, c in enumerate(df.columns):
            t.rows[0].cells[i].text = str(c)
        for _, r in df.iterrows():
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = f"{v:.3f}" if isinstance(v, float) else str(v)
        self.doc.add_paragraph()

    def references(self, keys):
        self.h("References (verify formatting and add your own sources before submission)")
        for k in keys:
            if k in REFS:
                self.doc.add_paragraph(REFS[k], style="List Number")

    def save(self, path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
        return path
